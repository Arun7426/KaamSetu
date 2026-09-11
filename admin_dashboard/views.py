from datetime import date, timedelta
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.models import User, Group
from django.contrib import messages
from django.db.models import Sum, Q, Avg, Count
from django.db import transaction
from django.core.exceptions import PermissionDenied
from django.utils import timezone
from datetime import datetime
from functools import wraps


from workers.models import Worker
from bookings.models import Booking, Review, Notification
from bookings.services import expire_pending_bookings
from payments.models import WorkerLedger, WorkerPaymentAlert, FeeSetting, Promotion

from io import BytesIO

from django.http import HttpResponse
from django.views.decorators.http import require_POST

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from .models import AuditLog, SystemSetting

from .audit_service import create_audit_log, get_client_ip

from django.contrib.auth import update_session_auth_hash
from django.contrib.auth.forms import PasswordChangeForm

# =========================================================
# EXISTING STAFF CHECK
# =========================================================

def staff_required(user):
    return user.is_authenticated and user.is_staff


# =========================================================
# KAAMSETU ADMIN ROLE
# =========================================================

def get_admin_role(user):
    """
    Two-level KaamSetu Admin architecture.

    SUPER_ADMIN:
        Django superuser

    ADMIN:
        Django staff user who is not a superuser

    NONE:
        Any other user
    """

    if not user or not user.is_authenticated:
        return "NONE"

    if user.is_superuser:
        return "SUPER_ADMIN"

    if user.is_staff:
        return "ADMIN"

    return "NONE"


# =========================================================
# KAAMSETU ADMIN ROLE FOUNDATION
# =========================================================

def admin_required(view_func):
    """
    Allow only KaamSetu Admin or Super Admin users.

    Super Admin:
        is_superuser = True

    Admin:
        is_staff = True
        is_superuser = False
    """

    @wraps(view_func)
    def wrapper(request, *args, **kwargs):

        if not request.user.is_authenticated:
            return redirect("login")

        # -------------------------------------------------
        # SUPER ADMIN
        # -------------------------------------------------

        if request.user.is_superuser:
            return view_func(
                request,
                *args,
                **kwargs
            )

        # -------------------------------------------------
        # NORMAL ADMIN
        # -------------------------------------------------

        if request.user.is_staff:
            return view_func(
                request,
                *args,
                **kwargs
            )

        # -------------------------------------------------
        # CUSTOMER / WORKER / NORMAL USER
        # -------------------------------------------------

        raise PermissionDenied(
            "You do not have permission to access the KaamSetu Admin Panel."
        )

    return wrapper


# =========================================================
# KAAMSETU PERMISSION REQUIRED
# =========================================================

def permission_required(permission_name):
    """
    Allow Super Admin automatically.

    Normal Admin users must have the
    specified Django permission.
    """

    def decorator(view_func):

        @wraps(view_func)
        def wrapper(request, *args, **kwargs):

            if not request.user.is_authenticated:
                return redirect("login")

            # -------------------------------------------------
            # SUPER ADMIN
            # -------------------------------------------------

            if request.user.is_superuser:
                return view_func(
                    request,
                    *args,
                    **kwargs
                )

            # -------------------------------------------------
            # NORMAL ADMIN
            # -------------------------------------------------

            if (
                request.user.is_staff
                and request.user.has_perm(permission_name)
            ):
                return view_func(
                    request,
                    *args,
                    **kwargs
                )

            raise PermissionDenied(
                "You do not have permission to access this section."
            )

        return wrapper

    return decorator

# =========================================================
# SUPER ADMIN ONLY
# =========================================================

def super_admin_required(view_func):
    """
    Allow ONLY Django Super Admin users.

    Normal Admin users are explicitly denied.
    """

    @wraps(view_func)
    def wrapper(request, *args, **kwargs):

        if not request.user.is_authenticated:
            return redirect("login")

        if not request.user.is_superuser:
            raise PermissionDenied(
                "Only Super Admin can access this section."
            )

        return view_func(
            request,
            *args,
            **kwargs
        )

    return wrapper

# =========================================================
# ADMIN DASHBOARD
# =========================================================

@admin_required
def admin_dashboard(request):

    role = get_admin_role(request.user)

    # =====================================================
    # USER COUNTS
    # =====================================================

    total_users = User.objects.count()

    total_workers = Worker.objects.count()

    total_customers = User.objects.filter(
        is_staff=False,
        worker_profile__isnull=True
    ).count()

    # =====================================================
    # BOOKING COUNTS
    # =====================================================

    total_bookings = Booking.objects.count()

    pending_bookings = Booking.objects.filter(
        status="Pending"
    ).count()

    accepted_bookings = Booking.objects.filter(
        status="Accepted"
    ).count()

    completed_bookings = Booking.objects.filter(
        status="Completed"
    ).count()

    cancelled_bookings = Booking.objects.filter(
        status="Cancelled"
    ).count()

    # =====================================================
    # FINANCE
    # SUPER ADMIN ONLY
    # =====================================================

    total_outstanding = None
    total_paid = None

    workers_with_outstanding = Worker.objects.none()
    recent_payments = WorkerLedger.objects.none()

    if request.user.is_superuser:

        total_outstanding = sum(
            entry.amount
            for entry in WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
                status="Pending"
            )
        )

        total_paid = sum(
            entry.amount
            for entry in WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
                status="Paid"
            )
        )

        workers_with_outstanding = Worker.objects.annotate(
            outstanding=Sum(
                "ledger_entries__amount",
                filter=Q(
                    ledger_entries__transaction_type="Booking Fee",
                    ledger_entries__status="Pending"
                )
            )
        ).filter(
            outstanding__gt=0
        ).order_by(
            "-outstanding"
        )

        recent_payments = WorkerLedger.objects.select_related(
            "worker",
            "booking"
        ).order_by(
            "-created_at"
        )[:5]

    # =====================================================
    # RECENT BOOKINGS
    # =====================================================

    recent_bookings = Booking.objects.select_related(
        "worker",
        "customer"
    ).order_by(
        "-created_at"
    )[:10]

    # =====================================================
    # RECENT CUSTOMERS
    # =====================================================

    recent_customers = User.objects.filter(
        is_staff=False,
        worker_profile__isnull=True
    ).order_by(
        "-date_joined"
    )[:5]

    # =====================================================
    # RECENT WORKERS
    # =====================================================

    recent_workers = Worker.objects.select_related(
        "user"
    ).order_by(
        "-user__date_joined"
    )[:5]

    return render(
        request,
        "admin_dashboard/dashboard.html",
        {
            "admin_role": role,

            # -------------------------------------------------
            # USER STATISTICS
            # -------------------------------------------------

            "total_users": total_users,
            "total_workers": total_workers,
            "total_customers": total_customers,

            # -------------------------------------------------
            # BOOKING STATISTICS
            # -------------------------------------------------

            "total_bookings": total_bookings,
            "pending_bookings": pending_bookings,
            "accepted_bookings": accepted_bookings,
            "completed_bookings": completed_bookings,
            "cancelled_bookings": cancelled_bookings,

            # -------------------------------------------------
            # FINANCE
            # SUPER ADMIN ONLY
            # -------------------------------------------------

            "total_outstanding": total_outstanding,
            "total_paid": total_paid,
            "workers_with_outstanding":
                workers_with_outstanding,
            "recent_payments":
                recent_payments,

            # -------------------------------------------------
            # RECENT ACTIVITY
            # -------------------------------------------------

            "recent_bookings":
                recent_bookings,
            "recent_customers":
                recent_customers,
            "recent_workers":
                recent_workers,
        }
    )


# =========================================================
# WORKERS
# =========================================================

@permission_required("workers.view_worker")
def admin_workers(request):

    workers = Worker.objects.select_related(
        "user"
    ).order_by(
        "name"
    )

    return render(
        request,
        "admin_dashboard/workers.html",
        {
            "workers": workers,
        }
    )


# =========================================================
# CUSTOMERS
# =========================================================

@permission_required("auth.view_user")
def admin_customers(request):

    customers = User.objects.filter(
        is_staff=False,
        worker_profile__isnull=True
    ).order_by(
        "username"
    )

    return render(
        request,
        "admin_dashboard/customers.html",
        {
            "customers": customers,
        }
    )


# =========================================================
# FRAUD CONTROL
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_fraud_control(request):
    """
    Central Super Admin-only fraud/safety control.

    Blocking uses Django's existing User.is_active flag so a blocked
    worker/customer cannot authenticate again. No new model or migration
    is required.

    Deleting a worker also deletes the linked User account and related
    cascade records according to the existing model relationships.
    """

    workers = Worker.objects.select_related(
        "user"
    ).order_by(
        "name"
    )

    customers = User.objects.filter(
        is_staff=False,
        worker_profile__isnull=True
    ).select_related(
        "customer_profile"
    ).order_by(
        "username"
    )

    return render(
        request,
        "admin_dashboard/fraud_control.html",
        {
            "workers": workers,
            "customers": customers,
        }
    )


@super_admin_required
@require_POST
def admin_toggle_worker_block(request, worker_id):
    """Block or unblock a worker. Super Admin only."""

    worker = get_object_or_404(
        Worker.objects.select_related("user"),
        id=worker_id
    )

    if worker.user is None:
        messages.error(
            request,
            f"Worker '{worker.name}' has no linked user account."
        )
        return redirect("admin_fraud_control")

    # A worker profile must never be allowed to control a staff account.
    if worker.user.is_staff or worker.user.is_superuser:
        messages.error(
            request,
            "This account cannot be managed through Fraud Control."
        )
        return redirect("admin_fraud_control")

    was_blocked = not worker.user.is_active
    worker.user.is_active = was_blocked
    worker.user.save(update_fields=["is_active"])

    action = "UPDATE"
    state = "unblocked" if was_blocked else "blocked"

    create_audit_log(
        admin=request.user,
        action=action,
        module="Fraud Control",
        description=(
            f"Worker account '{worker.name}' "
            f"(username: {worker.user.username}) was {state}."
        ),
        target_user=worker.user,
        target_id=worker.id,
        ip_address=get_client_ip(request),
    )

    messages.success(
        request,
        f"Worker '{worker.name}' {state} successfully."
    )

    return redirect("admin_fraud_control")


@super_admin_required
@require_POST
def admin_toggle_customer_block(request, user_id):
    """Block or unblock a customer. Super Admin only."""

    customer = get_object_or_404(
        User,
        id=user_id,
        is_staff=False,
        is_superuser=False,
        worker_profile__isnull=True,
    )

    was_blocked = not customer.is_active
    customer.is_active = was_blocked
    customer.save(update_fields=["is_active"])

    state = "unblocked" if was_blocked else "blocked"

    create_audit_log(
        admin=request.user,
        action="UPDATE",
        module="Fraud Control",
        description=(
            f"Customer account '{customer.username}' was {state}."
        ),
        target_user=customer,
        target_id=customer.id,
        ip_address=get_client_ip(request),
    )

    messages.success(
        request,
        f"Customer '{customer.username}' {state} successfully."
    )

    return redirect("admin_fraud_control")


@super_admin_required
@require_POST
def admin_delete_worker(request, worker_id):
    """Permanently delete a worker and its linked user account."""

    worker = get_object_or_404(
        Worker.objects.select_related("user"),
        id=worker_id
    )

    user = worker.user

    if user and (user.is_staff or user.is_superuser):
        messages.error(
            request,
            "This account cannot be deleted through Fraud Control."
        )
        return redirect("admin_fraud_control")

    worker_name = worker.name
    username = user.username if user else "No linked user"

    with transaction.atomic():

        create_audit_log(
            admin=request.user,
            action="DELETE",
            module="Fraud Control",
            description=(
                f"Worker account '{worker_name}' "
                f"(username: {username}) was permanently deleted "
                f"by Super Admin."
            ),
            target_user=user,
            target_id=worker.id,
            ip_address=get_client_ip(request),
        )

        if user:
            user.delete()
        else:
            worker.delete()

    messages.success(
        request,
        f"Worker '{worker_name}' deleted successfully."
    )

    return redirect("admin_fraud_control")


@super_admin_required
@require_POST
def admin_delete_customer(request, user_id):
    """Permanently delete a customer account."""

    customer = get_object_or_404(
        User,
        id=user_id,
        is_staff=False,
        is_superuser=False,
        worker_profile__isnull=True,
    )

    username = customer.username
    customer_id = customer.id

    with transaction.atomic():

        create_audit_log(
            admin=request.user,
            action="DELETE",
            module="Fraud Control",
            description=(
                f"Customer account '{username}' "
                f"was permanently deleted by Super Admin."
            ),
            target_user=customer,
            target_id=customer_id,
            ip_address=get_client_ip(request),
        )

        customer.delete()

    messages.success(
        request,
        f"Customer '{username}' deleted successfully."
    )

    return redirect("admin_fraud_control")


# =========================================================
# BOOKINGS
# =========================================================

@permission_required("bookings.view_booking")
def admin_bookings(request):

    bookings = Booking.objects.select_related(
        "worker",
        "customer"
    ).order_by(
        "-created_at"
    )

    status_filter = request.GET.get("status")

    if status_filter in [
        "Pending",
        "Accepted",
        "Completed",
        "Cancelled",
    ]:
        bookings = bookings.filter(
            status=status_filter
        )

    return render(
        request,
        "admin_dashboard/bookings.html",
        {
            "bookings": bookings,
            "status_filter": status_filter,
        }
    )

# =========================================================
# REVIEWS
# =========================================================

@permission_required("bookings.view_review")
def admin_reviews(request):

    # =====================================================
    # ALL REVIEWS
    # =====================================================

    reviews = Review.objects.select_related(
        "booking",
        "worker",
        "customer"
    ).order_by(
        "-created_at"
    )

    # =====================================================
    # RATING FILTER
    # =====================================================

    rating_filter = request.GET.get("rating")

    if rating_filter in [
        "1",
        "2",
        "3",
        "4",
        "5",
    ]:
        reviews = reviews.filter(
            rating=int(rating_filter)
        )

    # =====================================================
    # REVIEW STATISTICS
    # =====================================================

    review_stats = Review.objects.aggregate(
        total=Count("id"),
        average=Avg("rating"),
    )

    total_reviews = review_stats["total"] or 0
    average_rating = review_stats["average"] or 0

    five_star_reviews = Review.objects.filter(
        rating=5
    ).count()

    low_rating_reviews = Review.objects.filter(
        rating__lte=2
    ).count()

    # =====================================================
    # CONTEXT
    # =====================================================

    return render(
        request,
        "admin_dashboard/reviews.html",
        {
            "reviews": reviews,

            "total_reviews":
                total_reviews,

            "average_rating":
                round(average_rating, 1),

            "five_star_reviews":
                five_star_reviews,

            "low_rating_reviews":
                low_rating_reviews,

            "rating_filter":
                rating_filter,
        }
    )


# =========================================================
# NOTIFICATIONS
# =========================================================

@permission_required("bookings.view_notification")
def admin_notifications(request):

    notifications = Notification.objects.select_related(
        "recipient",
        "booking",
    ).order_by(
        "-created_at"
    )

    total_notifications = Notification.objects.count()

    unread_notifications = Notification.objects.filter(
        is_read=False
    ).count()

    read_notifications = Notification.objects.filter(
        is_read=True
    ).count()

    return render(
        request,
        "admin_dashboard/notifications.html",
        {
            "notifications": notifications,
            "total_notifications": total_notifications,
            "unread_notifications": unread_notifications,
            "read_notifications": read_notifications,
        }
    )

# =========================================================
# PAYMENTS
# =========================================================

@permission_required("payments.view_workerledger")
def admin_payments(request):

    ledger_entries = WorkerLedger.objects.select_related(
        "worker",
        "booking"
    ).order_by(
        "-created_at"
    )

    # -----------------------------------------------------
    # STATUS FILTER
    # -----------------------------------------------------

    status_filter = request.GET.get("status")

    if status_filter in [
        "Pending",
        "Paid",
    ]:
        ledger_entries = ledger_entries.filter(
            status=status_filter
        )

    # -----------------------------------------------------
    # TRANSACTION FILTER
    # -----------------------------------------------------

    transaction_filter = request.GET.get("transaction")

    if transaction_filter in [
        "Booking Fee",
        "Payment",
    ]:
        ledger_entries = ledger_entries.filter(
            transaction_type=transaction_filter
        )

    return render(
        request,
        "admin_dashboard/payments.html",
        {
            "ledger_entries": ledger_entries,
            "status_filter": status_filter,
            "transaction_filter": transaction_filter,
        }
    )

# =========================================================
# FEE SETTINGS
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_fee_settings(request):
    """
    Manage the single active platform fee configuration.

    Fee Settings are intentionally restricted to Super Admin.
    Existing FeeSetting model is reused; no migration is required.
    """

    fee_setting = FeeSetting.objects.first()

    # Create the default only when the table has no setting at all.
    # An intentionally inactive setting must remain inactive.
    if fee_setting is None:
        fee_setting = FeeSetting.objects.create(
            fee_type="fixed",
            fee_value=Decimal("20.00"),
            is_active=True,
        )

    if request.method == "POST":

        fee_type = request.POST.get("fee_type")
        fee_value_raw = request.POST.get("fee_value", "").strip()
        is_active = request.POST.get("is_active") == "on"

        if fee_type not in {"fixed", "percentage"}:
            messages.error(request, "Please select a valid fee type.")
            return redirect("admin_fee_settings")

        try:
            fee_value = Decimal(fee_value_raw)
        except Exception:
            messages.error(request, "Please enter a valid fee value.")
            return redirect("admin_fee_settings")

        if fee_value < 0:
            messages.error(request, "Fee value cannot be negative.")
            return redirect("admin_fee_settings")

        if fee_type == "percentage" and fee_value > 100:
            messages.error(request, "Percentage fee cannot exceed 100%.")
            return redirect("admin_fee_settings")

        fee_setting.fee_type = fee_type
        fee_setting.fee_value = fee_value
        fee_setting.is_active = is_active
        fee_setting.save()
        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="Fee Settings",
            description=(
                f"Platform fee updated to "
                f"{fee_setting.fee_type} - "
                f"{fee_setting.fee_value}"
            ),
            target_id=fee_setting.id,
            ip_address=get_client_ip(request),
        )

        messages.success(request, "Fee Settings updated successfully.")
        return redirect("admin_fee_settings")

    return render(
        request,
        "admin_dashboard/fee_settings.html",
        {"fee_setting": fee_setting},
    )



# =========================================================
# PROMOTION MANAGEMENT
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_promotions(request):
    """
    Manage KaamSetu promotions.

    Current Promotion model is intentionally reused without adding
    fields or migrations. Only Super Admin can create, edit, activate,
    deactivate, or delete promotions.

    The payment service currently uses the newest active promotion,
    so this view keeps at most one promotion active at a time.
    """

    promotions = Promotion.objects.order_by(
        "-is_active",
        "-updated_at",
        "-created_at",
    )

    return render(
        request,
        "admin_dashboard/promotions.html",
        {
            "promotions": promotions,
        }
    )


@super_admin_required
def admin_promotion_create(request):
    """Create a new promotion. Super Admin only."""

    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        free_bookings_limit_raw = request.POST.get(
            "free_bookings_limit",
            ""
        ).strip()
        is_active = request.POST.get("is_active") == "on"

        if not name:
            messages.error(
                request,
                "Promotion name is required."
            )
            return redirect("admin_promotion_create")

        try:
            free_bookings_limit = int(free_bookings_limit_raw)
        except (TypeError, ValueError):
            messages.error(
                request,
                "Please enter a valid free booking limit."
            )
            return redirect("admin_promotion_create")

        if free_bookings_limit < 0:
            messages.error(
                request,
                "Free booking limit cannot be negative."
            )
            return redirect("admin_promotion_create")

        promotion = Promotion.objects.create(
            name=name,
            free_bookings_limit=free_bookings_limit,
            is_active=is_active,
        )

        # The current payment service selects the newest active
        # promotion. Keep only one active promotion at a time.
        if is_active:
            Promotion.objects.exclude(
                pk=promotion.pk
            ).update(
                is_active=False
            )
        create_audit_log(
            admin=request.user,
            action="CREATE",
            module="Promotions",
            description=(
                f'Promotion "{promotion.name}" created '
                f'with {promotion.free_bookings_limit} free bookings '
                f'and status '
                f'{"Active" if promotion.is_active else "Inactive"}.'
            ),
            target_id=promotion.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Promotion created successfully."
        )

        return redirect("admin_promotions")

    return render(
        request,
        "admin_dashboard/promotion_form.html",
        {
            "page_title": "Create Promotion",
            "form_action": "admin_promotion_create",
            "promotion": None,
        }
    )


@super_admin_required
def admin_promotion_edit(request, promotion_id):
    """Edit an existing promotion. Super Admin only."""

    promotion = Promotion.objects.filter(
        id=promotion_id
    ).first()

    if promotion is None:
        raise PermissionDenied("Promotion not found.")

    if request.method == "POST":
        name = request.POST.get("name", "").strip()
        free_bookings_limit_raw = request.POST.get(
            "free_bookings_limit",
            ""
        ).strip()
        is_active = request.POST.get("is_active") == "on"

        if not name:
            messages.error(
                request,
                "Promotion name is required."
            )
            return redirect(
                "admin_promotion_edit",
                promotion_id=promotion.id
            )

        try:
            free_bookings_limit = int(free_bookings_limit_raw)
        except (TypeError, ValueError):
            messages.error(
                request,
                "Please enter a valid free booking limit."
            )
            return redirect(
                "admin_promotion_edit",
                promotion_id=promotion.id
            )

        if free_bookings_limit < 0:
            messages.error(
                request,
                "Free booking limit cannot be negative."
            )
            return redirect(
                "admin_promotion_edit",
                promotion_id=promotion.id
            )

        promotion.name = name
        promotion.free_bookings_limit = free_bookings_limit
        promotion.is_active = is_active
        promotion.save()

        if is_active:
            Promotion.objects.exclude(
                pk=promotion.pk
            ).update(
                is_active=False
            )
        
        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="Promotions",
            description=(
                f'Promotion "{promotion.name}" updated '
                f'to {promotion.free_bookings_limit} free bookings '
                f'and status '
                f'{"Active" if promotion.is_active else "Inactive"}.'
            ),
            target_id=promotion.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Promotion updated successfully."
        )

        return redirect("admin_promotions")

    return render(
        request,
        "admin_dashboard/promotion_form.html",
        {
            "page_title": "Edit Promotion",
            "form_action": "admin_promotion_edit",
            "promotion": promotion,
        }
    )


@super_admin_required
def admin_promotion_toggle(request, promotion_id):
    """Activate/deactivate a promotion. Super Admin only."""

    if request.method != "POST":
        return redirect("admin_promotions")

    promotion = Promotion.objects.filter(
        id=promotion_id
    ).first()

    if promotion is None:
        raise PermissionDenied("Promotion not found.")

    promotion.is_active = not promotion.is_active
    promotion.save(update_fields=["is_active", "updated_at"])
    
    create_audit_log(
        admin=request.user,
        action="UPDATE",
        module="Promotions",
        description=(
            f'Promotion "{promotion.name}" '
            f'{"activated" if promotion.is_active else "deactivated"}.'
        ),
        target_id=promotion.id,
        ip_address=get_client_ip(request),
    )

    if promotion.is_active:
        Promotion.objects.exclude(
            pk=promotion.pk
        ).update(
            is_active=False
        )
        messages.success(
            request,
            f'Promotion "{promotion.name}" is now active.'
        )
    else:
        messages.success(
            request,
            f'Promotion "{promotion.name}" is now inactive.'
        )

    return redirect("admin_promotions")


@super_admin_required
def admin_promotion_delete(request, promotion_id):
    """Delete a promotion. Super Admin only."""

    promotion = Promotion.objects.filter(
        id=promotion_id
    ).first()

    if promotion is None:
        raise PermissionDenied("Promotion not found.")

    if request.method == "POST":
        
        name = promotion.name

        create_audit_log(
            admin=request.user,
            action="DELETE",
            module="Promotions",
            description=(
                f'Promotion "{name}" deleted.'
            ),
            target_id=promotion.id,
            ip_address=get_client_ip(request),
        )

        promotion.delete()

        messages.success(
            request,
            f'Promotion "{name}" deleted successfully.'
        )

        return redirect("admin_promotions")

    return render(
        request,
        "admin_dashboard/promotion_delete.html",
        {
            "promotion": promotion,
        }
    )


# =========================================================
# PAYMENT ALERTS
# SUPER ADMIN ONLY
# =========================================================

def super_admin_only(view_func):
    """
    Allow ONLY Django Super Admin users.
    """

    @wraps(view_func)
    def wrapper(request, *args, **kwargs):

        if not request.user.is_authenticated:
            return redirect("login")

        if not request.user.is_superuser:
            raise PermissionDenied(
                "Only Super Admin can access Payment Alerts."
            )

        return view_func(
            request,
            *args,
            **kwargs
        )

    return wrapper


@super_admin_only
def admin_payment_alerts(request):

    alerts = WorkerPaymentAlert.objects.select_related(
        "worker"
    ).order_by(
        "-created_at"
    )

    # -----------------------------------------------------
    # STATISTICS
    # -----------------------------------------------------

    total_alerts = WorkerPaymentAlert.objects.count()

    pending_followups = WorkerPaymentAlert.objects.filter(
        status__in=[
            "Pending",
            "Follow-up",
        ]
    ).count()

    sms_pending = WorkerPaymentAlert.objects.filter(
        sms_status="Pending"
    ).count()

    resolved_alerts = WorkerPaymentAlert.objects.filter(
        status="Resolved"
    ).count()

    # -----------------------------------------------------
    # CONTEXT
    # -----------------------------------------------------

    return render(
        request,
        "admin_dashboard/payment_alerts.html",
        {
            "alerts": alerts,

            "total_alerts":
                total_alerts,

            "pending_followups":
                pending_followups,

            "sms_pending":
                sms_pending,

            "resolved_alerts":
                resolved_alerts,
        }
    )
    

# =========================================================
# PAYMENT ALERT DETAILS
# SUPER ADMIN ONLY
# =========================================================

@super_admin_only
def admin_payment_alert_detail(request, alert_id):

    alert = WorkerPaymentAlert.objects.select_related(
        "worker"
    ).get(
        id=alert_id
    )

    worker = alert.worker

    # -----------------------------------------------------
    # PENDING LEDGER ENTRIES
    # -----------------------------------------------------

    pending_ledger_entries = WorkerLedger.objects.filter(
        worker=worker,
        transaction_type="Booking Fee",
        status="Pending",
    ).select_related(
        "booking"
    ).order_by(
        "created_at"
    )

    # -----------------------------------------------------
    # CURRENT OUTSTANDING
    # -----------------------------------------------------

    outstanding_result = WorkerLedger.objects.filter(
        worker=worker,
        transaction_type="Booking Fee",
        status="Pending",
    ).aggregate(
        total=Sum("amount")
    )

    current_outstanding = (
        outstanding_result["total"] or 0
    )

    # -----------------------------------------------------
    # LAST BOOKING
    # -----------------------------------------------------

    last_booking = Booking.objects.filter(
        worker=worker
    ).order_by(
        "-created_at"
    ).first()

    # -----------------------------------------------------
    # LAST PAYMENT
    # -----------------------------------------------------

    last_payment = WorkerLedger.objects.filter(
        worker=worker,
        transaction_type="Payment",
        status="Paid",
    ).order_by(
        "-paid_at",
        "-created_at"
    ).first()

    # -----------------------------------------------------
    # ALERT HISTORY
    # -----------------------------------------------------

    alert_history = WorkerPaymentAlert.objects.filter(
        worker=worker
    ).order_by(
        "-created_at"
    )

    return render(
        request,
        "admin_dashboard/payment_alert_detail.html",
        {
            "alert": alert,
            "worker": worker,

            "pending_ledger_entries":
                pending_ledger_entries,

            "current_outstanding":
                current_outstanding,

            "last_booking":
                last_booking,

            "last_payment":
                last_payment,

            "alert_history":
                alert_history,
        }
    )
# =========================================================
# MARK PAYMENT ALERT FOLLOW-UP
# SUPER ADMIN ONLY
# =========================================================

@super_admin_only
def admin_payment_alert_follow_up(request, alert_id):

    if request.method != "POST":
        raise PermissionDenied(
            "Follow-up action must use POST."
        )

    alert = WorkerPaymentAlert.objects.get(
        id=alert_id
    )

    from payments.alert_service import mark_alert_follow_up

    mark_alert_follow_up(alert)

    return redirect(
        "admin_payment_alert_detail",
        alert_id=alert.id
    )


# =========================================================
# SEND PAYMENT ALERT SMS
# SUPER ADMIN ONLY
# =========================================================

@super_admin_only
def admin_payment_alert_send_sms(request, alert_id):

    if request.method != "POST":
        raise PermissionDenied(
            "SMS action must use POST."
        )

    alert = WorkerPaymentAlert.objects.get(
        id=alert_id
    )

    from payments.sms_service import send_alert_sms

    send_alert_sms(alert)

    return redirect(
        "admin_payment_alert_detail",
        alert_id=alert.id
    )

    
# =========================================================
# WORKER LEDGER
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_worker_ledger(request):

    ledger_entries = WorkerLedger.objects.select_related(
        "worker",
        "booking"
    ).order_by(
        "-created_at"
    )

    status_filter = request.GET.get("status")

    if status_filter in [
        "Pending",
        "Paid",
    ]:
        ledger_entries = ledger_entries.filter(
            status=status_filter
        )

    transaction_filter = request.GET.get("transaction")

    if transaction_filter in [
        "Booking Fee",
        "Payment",
    ]:
        ledger_entries = ledger_entries.filter(
            transaction_type=transaction_filter
        )

    # =====================================================
    # LEDGER SUMMARY
    # =====================================================

    total_ledger_amount = (
        WorkerLedger.objects.aggregate(
            total=Sum("amount")
        )["total"] or 0
    )

    pending_amount = (
        WorkerLedger.objects.filter(
            status="Pending"
        ).aggregate(
            total=Sum("amount")
        )["total"] or 0
    )

    paid_amount = (
        WorkerLedger.objects.filter(
            status="Paid"
        ).aggregate(
            total=Sum("amount")
        )["total"] or 0
    )

    return render(
        request,
        "admin_dashboard/worker_ledger.html",
        {
            "ledger_entries": ledger_entries,
            "status_filter": status_filter,
            "transaction_filter": transaction_filter,

            "total_ledger_amount":
                total_ledger_amount,

            "pending_amount":
                pending_amount,

            "paid_amount":
                paid_amount,
        }
    )


# =========================================================
# INSIGHTS & ANALYTICS
# =========================================================

@permission_required("bookings.view_booking")
def admin_insights(request):
    """
    KaamSetu Insights & Analytics.

    Supported periods:
        - daily
        - weekly
        - monthly
        - yearly
        - custom

    The selected period controls the operational and financial
    analytics shown on the Insights page. Current outstanding
    remains a live/global figure because it is not a historical
    period metric.

    Finance figures are intentionally available only to Super Admin.
    """

    # =====================================================
    # PERIOD / DATE FILTER
    # =====================================================

    today = timezone.localdate()

    period = request.GET.get("period", "monthly").lower()
    if period not in [
        "daily",
        "weekly",
        "monthly",
        "yearly",
        "custom",
    ]:
        period = "monthly"

    start_date = None
    end_date = today
    period_label = "Monthly"
    chart_granularity = "month"
    custom_start = request.GET.get("start_date", "")
    custom_end = request.GET.get("end_date", "")

    def subtract_months(year, month, count):
        total = (year * 12 + (month - 1)) - count
        return total // 12, (total % 12) + 1

    if period == "daily":
        # Last 30 days, shown day-by-day.
        start_date = today - timedelta(days=29)
        period_label = "Daily — Last 30 Days"
        chart_granularity = "day"

    elif period == "weekly":
        # Last 12 weeks, shown week-by-week.
        start_date = today - timedelta(days=83)
        period_label = "Weekly — Last 12 Weeks"
        chart_granularity = "week"

    elif period == "monthly":
        # Current month + previous 11 months.
        start_year, start_month = subtract_months(
            today.year,
            today.month,
            11
        )
        start_date = date(start_year, start_month, 1)
        period_label = "Monthly — Last 12 Months"
        chart_granularity = "month"

    elif period == "yearly":
        # Current year + previous 4 years.
        start_date = date(today.year - 4, 1, 1)
        period_label = "Yearly — Last 5 Years"
        chart_granularity = "year"

    else:
        # -------------------------------------------------
        # CUSTOM DATE RANGE
        # -------------------------------------------------
        try:
            parsed_start = date.fromisoformat(custom_start)
            parsed_end = date.fromisoformat(custom_end)

            if parsed_start > parsed_end:
                parsed_start, parsed_end = parsed_end, parsed_start

            start_date = parsed_start
            end_date = parsed_end
            period_label = (
                f"Custom — {start_date.strftime('%d %b %Y')}"
                f" to {end_date.strftime('%d %b %Y')}"
            )

        except (TypeError, ValueError):
            # Safe fallback if custom dates are missing/invalid.
            period = "monthly"
            start_year, start_month = subtract_months(
                today.year,
                today.month,
                11
            )
            start_date = date(start_year, start_month, 1)
            end_date = today
            period_label = "Monthly — Last 12 Months"
            chart_granularity = "month"

        if period == "custom":
            span_days = (end_date - start_date).days + 1

            # Keep custom charts readable automatically.
            if span_days <= 31:
                chart_granularity = "day"
            elif span_days <= 180:
                chart_granularity = "week"
            else:
                chart_granularity = "month"

    # =====================================================
    # BASE QUERYSETS FOR SELECTED PERIOD
    # =====================================================

    period_bookings = Booking.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
    )

    period_users = User.objects.filter(
        date_joined__date__gte=start_date,
        date_joined__date__lte=end_date,
    )

    period_reviews = Review.objects.filter(
        created_at__date__gte=start_date,
        created_at__date__lte=end_date,
    )

    # =====================================================
    # GLOBAL USER TOTALS
    # =====================================================

    total_users = User.objects.count()
    total_workers = Worker.objects.count()
    total_customers = User.objects.filter(
        is_staff=False,
        worker_profile__isnull=True,
    ).count()

    # =====================================================
    # PERIOD USER GROWTH
    # =====================================================

    new_users = period_users.count()

    new_workers = period_users.filter(
        worker_profile__isnull=False,
    ).count()

    new_customers = period_users.filter(
        is_staff=False,
        worker_profile__isnull=True,
    ).count()

    # =====================================================
    # PERIOD BOOKING STATISTICS
    # =====================================================

    total_bookings = period_bookings.count()

    pending_bookings = period_bookings.filter(
        status="Pending"
    ).count()

    accepted_bookings = period_bookings.filter(
        status="Accepted"
    ).count()

    completed_bookings = period_bookings.filter(
        status="Completed"
    ).count()

    cancelled_bookings = period_bookings.filter(
        status="Cancelled"
    ).count()

    completion_rate = round(
        (completed_bookings / total_bookings) * 100,
        1
    ) if total_bookings else 0

    cancellation_rate = round(
        (cancelled_bookings / total_bookings) * 100,
        1
    ) if total_bookings else 0

    # =====================================================
    # PERIOD REVIEW / RATING STATISTICS
    # =====================================================

    review_stats = period_reviews.aggregate(
        total=Count("id"),
        average=Avg("rating"),
    )

    total_reviews = review_stats["total"] or 0
    average_rating = review_stats["average"] or 0
    average_rating = round(float(average_rating), 1)

    five_star_reviews = period_reviews.filter(
        rating=5
    ).count()

    low_rating_reviews = period_reviews.filter(
        rating__lte=2
    ).count()

    # =====================================================
    # TREND BUCKETS
    # =====================================================

    trend_labels = []
    trend_booking_counts = []
    trend_pending_counts = []
    trend_accepted_counts = []
    trend_completed_counts = []
    trend_cancelled_counts = []
    trend_user_counts = []

    # Finance trend data (Super Admin only)
    finance_platform_fee_counts = []
    finance_received_payment_counts = []

    buckets = []

    if chart_granularity == "day":
        current = start_date
        while current <= end_date:
            buckets.append((current, current))
            current += timedelta(days=1)

    elif chart_granularity == "week":
        current = start_date - timedelta(
            days=start_date.weekday()
        )
        last = end_date

        while current <= last:
            bucket_end = current + timedelta(days=6)
            buckets.append((current, bucket_end))
            current += timedelta(days=7)

    elif chart_granularity == "month":
        current_year = start_date.year
        current_month = start_date.month

        while True:
            bucket_start = date(
                current_year,
                current_month,
                1
            )

            if current_month == 12:
                next_month = date(
                    current_year + 1,
                    1,
                    1
                )
            else:
                next_month = date(
                    current_year,
                    current_month + 1,
                    1
                )

            bucket_end = next_month - timedelta(days=1)
            buckets.append((bucket_start, bucket_end))

            if bucket_end >= end_date:
                break

            if current_month == 12:
                current_month = 1
                current_year += 1
            else:
                current_month += 1

    else:
        for current_year in range(
            start_date.year,
            end_date.year + 1
        ):
            bucket_start = date(current_year, 1, 1)
            bucket_end = date(current_year, 12, 31)
            buckets.append((bucket_start, bucket_end))

    for bucket_start, bucket_end in buckets:
        actual_start = max(bucket_start, start_date)
        actual_end = min(bucket_end, end_date)

        bucket_bookings = Booking.objects.filter(
            created_at__date__gte=actual_start,
            created_at__date__lte=actual_end,
        )

        bucket_users = User.objects.filter(
            date_joined__date__gte=actual_start,
            date_joined__date__lte=actual_end,
        )

        if chart_granularity == "day":
            label = bucket_start.strftime("%d %b")
        elif chart_granularity == "week":
            label = (
                f"{actual_start.strftime('%d %b')}"
                f" – {actual_end.strftime('%d %b')}"
            )
        elif chart_granularity == "month":
            label = bucket_start.strftime("%b %Y")
        else:
            label = bucket_start.strftime("%Y")

        trend_labels.append(label)
        trend_booking_counts.append(bucket_bookings.count())
        trend_pending_counts.append(
            bucket_bookings.filter(status="Pending").count()
        )
        trend_accepted_counts.append(
            bucket_bookings.filter(status="Accepted").count()
        )
        trend_completed_counts.append(
            bucket_bookings.filter(status="Completed").count()
        )
        trend_cancelled_counts.append(
            bucket_bookings.filter(status="Cancelled").count()
        )
        trend_user_counts.append(bucket_users.count())

        # -------------------------------------------------
        # FINANCE TREND (SUPER ADMIN ONLY)
        # -------------------------------------------------
        if request.user.is_superuser:
            bucket_platform_fees = WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
                created_at__date__gte=actual_start,
                created_at__date__lte=actual_end,
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")

            bucket_received_payments = WorkerLedger.objects.filter(
                transaction_type="Payment",
                status="Paid",
                paid_at__date__gte=actual_start,
                paid_at__date__lte=actual_end,
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")

            finance_platform_fee_counts.append(
                float(bucket_platform_fees)
            )
            finance_received_payment_counts.append(
                float(bucket_received_payments)
            )

    # =====================================================
    # BOOKING STATUS DISTRIBUTION
    # =====================================================

    status_labels = [
        "Pending",
        "Accepted",
        "Completed",
        "Cancelled",
    ]

    status_counts = [
        pending_bookings,
        accepted_bookings,
        completed_bookings,
        cancelled_bookings,
    ]

    # =====================================================
    # FINANCE
    # SUPER ADMIN ONLY
    # =====================================================

    total_outstanding = None
    total_paid = None
    total_platform_fees = None

    period_platform_fees = None
    period_received_payments = None
    period_pending_fees = None
    period_paid_fees = None
    workers_with_outstanding = Worker.objects.none()
    period_ledger_entries = WorkerLedger.objects.none()

    if request.user.is_superuser:
        # -------------------------------------------------
        # LIVE / GLOBAL FINANCE
        # -------------------------------------------------

        total_outstanding = (
            WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
                status="Pending",
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        total_paid = (
            WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
                status="Paid",
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        total_platform_fees = (
            WorkerLedger.objects.filter(
                transaction_type="Booking Fee",
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        # -------------------------------------------------
        # SELECTED PERIOD FINANCE
        # -------------------------------------------------

        period_booking_fees = WorkerLedger.objects.filter(
            transaction_type="Booking Fee",
            created_at__date__gte=start_date,
            created_at__date__lte=end_date,
        )

        period_platform_fees = (
            period_booking_fees.aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        period_pending_fees = (
            period_booking_fees.filter(
                status="Pending"
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        period_paid_fees = (
            period_booking_fees.filter(
                status="Paid"
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        # Payment transactions represent actual worker payments received.
        period_received_payments = (
            WorkerLedger.objects.filter(
                transaction_type="Payment",
                status="Paid",
                paid_at__date__gte=start_date,
                paid_at__date__lte=end_date,
            ).aggregate(
                total=Sum("amount")
            )["total"] or Decimal("0.00")
        )

        # -------------------------------------------------
        # WORKER-WISE LIVE OUTSTANDING
        # -------------------------------------------------

        workers_with_outstanding = Worker.objects.annotate(
            outstanding=Sum(
                "ledger_entries__amount",
                filter=Q(
                    ledger_entries__transaction_type="Booking Fee",
                    ledger_entries__status="Pending",
                )
            )
        ).filter(
            outstanding__gt=0
        ).order_by(
            "-outstanding",
            "name",
        )

        # -------------------------------------------------
        # PERIOD LEDGER ACTIVITY
        # -------------------------------------------------

        period_ledger_entries = WorkerLedger.objects.filter(
            created_at__date__gte=start_date,
            created_at__date__lte=end_date,
        ).select_related(
            "worker",
            "booking",
        ).order_by(
            "-created_at"
        )[:20]

    # =====================================================
    # CONTEXT
    # =====================================================

    context = {
        # ---------------------------------------------
        # PERIOD / FILTER
        # ---------------------------------------------
        "selected_period": period,
        "period_label": period_label,
        "start_date": start_date,
        "end_date": end_date,
        "custom_start": custom_start,
        "custom_end": custom_end,

        # ---------------------------------------------
        # GLOBAL USER TOTALS
        # ---------------------------------------------
        "total_users": total_users,
        "total_workers": total_workers,
        "total_customers": total_customers,

        # ---------------------------------------------
        # PERIOD USER GROWTH
        # ---------------------------------------------
        "new_users": new_users,
        "new_workers": new_workers,
        "new_customers": new_customers,

        # ---------------------------------------------
        # PERIOD BOOKING STATISTICS
        # ---------------------------------------------
        "total_bookings": total_bookings,
        "pending_bookings": pending_bookings,
        "accepted_bookings": accepted_bookings,
        "completed_bookings": completed_bookings,
        "cancelled_bookings": cancelled_bookings,
        "completion_rate": completion_rate,
        "cancellation_rate": cancellation_rate,

        # ---------------------------------------------
        # PERIOD REVIEWS
        # ---------------------------------------------
        "total_reviews": total_reviews,
        "average_rating": average_rating,
        "five_star_reviews": five_star_reviews,
        "low_rating_reviews": low_rating_reviews,

        # ---------------------------------------------
        # MAIN TREND CHART
        # ---------------------------------------------
        "chart_granularity": chart_granularity,
        "trend_labels": trend_labels,
        "trend_booking_counts": trend_booking_counts,
        "trend_pending_counts": trend_pending_counts,
        "trend_accepted_counts": trend_accepted_counts,
        "trend_completed_counts": trend_completed_counts,
        "trend_cancelled_counts": trend_cancelled_counts,
        "trend_user_counts": trend_user_counts,

        # ---------------------------------------------
        # FINANCE TREND CHART
        # ---------------------------------------------
        "finance_platform_fee_counts": finance_platform_fee_counts,
        "finance_received_payment_counts": finance_received_payment_counts,

        # ---------------------------------------------
        # STATUS CHART
        # ---------------------------------------------
        "status_labels": status_labels,
        "status_counts": status_counts,

        # ---------------------------------------------
        # FINANCE - SUPER ADMIN ONLY
        # ---------------------------------------------
        "total_outstanding": total_outstanding,
        "total_paid": total_paid,
        "total_platform_fees": total_platform_fees,
        "period_platform_fees": period_platform_fees,
        "period_received_payments": period_received_payments,
        "period_pending_fees": period_pending_fees,
        "period_paid_fees": period_paid_fees,
        "workers_with_outstanding": workers_with_outstanding,
        "period_ledger_entries": period_ledger_entries,
    }

    return render(
        request,
        "admin_dashboard/insights.html",
        context,
    )


# =========================================================
# REPORTS
# =========================================================

@admin_required
def admin_reports(request):
    """
    KaamSetu Report Generator.

    Reports are generated according to:
        1. Report type
        2. Selected date range

    Export functionality will use the same filtered data.
    Financial reports remain Super Admin only.
    """

    today = timezone.localdate()

    # =====================================================
    # REPORT TYPE
    # =====================================================

    report_type = request.GET.get(
        "report_type",
        "new_users"
    )

    allowed_reports = [
        "new_users",
        "active_users",
        "inactive_users",
        "bookings",
        "accepted_bookings",
        "completed_bookings",
        "cancelled_bookings",
        "pending_bookings",
        "new_workers",
        "active_workers",
        "worker_performance",
        "reviews",
        "low_ratings",
        "platform_fees",
        "worker_payments",
        "outstanding",
        "ledger",
        "app_retention",
        "audit_logs",
    ]

    if report_type not in allowed_reports:
        report_type = "new_users"

    # =====================================================
    # DATE FILTER
    # =====================================================

    period = request.GET.get(
        "period",
        "30"
    )

    start_date = None
    end_date = today

    custom_start = request.GET.get(
        "start_date",
        ""
    )

    custom_end = request.GET.get(
        "end_date",
        ""
    )

    period_label = "Last 30 Days"

    # -----------------------------------------------------
    # ALL TIME
    # -----------------------------------------------------

    if period == "all":

        start_date = None
        period_label = "All Time"

    # -----------------------------------------------------
    # TODAY
    # -----------------------------------------------------

    elif period == "today":

        start_date = today
        period_label = "Today"

    # -----------------------------------------------------
    # LAST 7 DAYS
    # -----------------------------------------------------

    elif period == "7":

        start_date = today - timedelta(days=6)
        period_label = "Last 7 Days"

    # -----------------------------------------------------
    # LAST 30 DAYS
    # -----------------------------------------------------

    elif period == "30":

        start_date = today - timedelta(days=29)
        period_label = "Last 30 Days"

    # -----------------------------------------------------
    # THIS MONTH
    # -----------------------------------------------------

    elif period == "month":

        start_date = today.replace(day=1)
        period_label = "This Month"

    # -----------------------------------------------------
    # LAST MONTH
    # -----------------------------------------------------

    elif period == "last_month":

        first_this_month = today.replace(day=1)

        last_month_end = (
            first_this_month - timedelta(days=1)
        )

        start_date = last_month_end.replace(day=1)
        end_date = last_month_end

        period_label = "Last Month"

    # -----------------------------------------------------
    # THIS YEAR
    # -----------------------------------------------------

    elif period == "year":

        start_date = date(
            today.year,
            1,
            1
        )

        period_label = "This Year"

    # -----------------------------------------------------
    # CUSTOM
    # -----------------------------------------------------

    elif period == "custom":

        try:

            parsed_start = date.fromisoformat(
                custom_start
            )

            parsed_end = date.fromisoformat(
                custom_end
            )

            if parsed_start > parsed_end:

                parsed_start, parsed_end = (
                    parsed_end,
                    parsed_start
                )

            start_date = parsed_start
            end_date = parsed_end

            period_label = (
                f"{start_date.strftime('%d %b %Y')}"
                f" – "
                f"{end_date.strftime('%d %b %Y')}"
            )

        except (
            TypeError,
            ValueError
        ):

            period = "30"

            start_date = (
                today -
                timedelta(days=29)
            )

            end_date = today

            period_label = "Last 30 Days"

    # =====================================================
    # REPORT DEFINITIONS
    # =====================================================

    report_definitions = {

        # -------------------------------------------------
        # USERS
        # -------------------------------------------------

        "new_users": {
            "label": "New Users",
            "category": "Users",
            "description":
                "Users registered during the selected period.",
        },

        "active_users": {
            "label": "Active Users",
            "category": "Users",
            "description":
                "Users who were active during the selected period.",
        },

        "inactive_users": {
            "label": "Inactive Users",
            "category": "Users",
            "description":
                "Users with no recent login activity.",
        },

        # -------------------------------------------------
        # BOOKINGS
        # -------------------------------------------------

        "bookings": {
            "label": "All Bookings",
            "category": "Bookings",
            "description":
                "All bookings created during the selected period.",
        },

        "accepted_bookings": {
            "label": "Accepted Bookings",
            "category": "Bookings",
            "description":
                "Bookings accepted during the selected period.",
        },

        "completed_bookings": {
            "label": "Completed Bookings",
            "category": "Bookings",
            "description":
                "Bookings completed during the selected period.",
        },

        "cancelled_bookings": {
            "label": "Cancelled Bookings",
            "category": "Bookings",
            "description":
                "Bookings cancelled during the selected period.",
        },

        "pending_bookings": {
            "label": "Pending Bookings",
            "category": "Bookings",
            "description":
                "Bookings currently pending.",
        },

        # -------------------------------------------------
        # WORKERS
        # -------------------------------------------------

        "new_workers": {
            "label": "New Workers",
            "category": "Workers",
            "description":
                "Workers registered during the selected period.",
        },

        "active_workers": {
            "label": "Active Workers",
            "category": "Workers",
            "description":
                "Workers with recent platform activity.",
        },

        "worker_performance": {
            "label": "Worker Performance",
            "category": "Workers",
            "description":
                "Worker-wise booking performance.",
        },

        # -------------------------------------------------
        # REVIEWS
        # -------------------------------------------------

        "reviews": {
            "label": "Reviews",
            "category": "Reviews",
            "description":
                "Reviews submitted during the selected period.",
        },

        "low_ratings": {
            "label": "Low Ratings",
            "category": "Reviews",
            "description":
                "Reviews with ratings of 2 stars or below.",
        },

        # -------------------------------------------------
        # FINANCE
        # -------------------------------------------------

        "platform_fees": {
            "label": "Platform Fee Report",
            "category": "Finance",
            "description":
                "Platform fees generated during the selected period.",
        },

        "worker_payments": {
            "label": "Worker Payment Report",
            "category": "Finance",
            "description":
                "Worker payments recorded during the selected period.",
        },

        "outstanding": {
            "label": "Outstanding Report",
            "category": "Finance",
            "description":
                "Current worker outstanding balances.",
        },

        "ledger": {
            "label": "Worker Ledger Report",
            "category": "Finance",
            "description":
                "Worker ledger transactions during the selected period.",
        },

        # -------------------------------------------------
        # APP RETENTION
        # -------------------------------------------------

        "app_retention": {
            "label": "App Retention",
            "category": "App & Retention",
            "description":
                "App installation, activity and retention analytics.",
        },
        
        # -------------------------------------------------
        # AUDIT LOGS
        # -------------------------------------------------

        "audit_logs": {
            "label": "Audit Logs",
            "category": "System",
            "description":
                "Administrative activity and system audit history.",
        },
    }

    selected_report = report_definitions[
        report_type
    ]

    # =====================================================
    # INITIAL DATA
    # =====================================================

    report_rows = []

    report_columns = []

    report_count = 0

    report_total = Decimal("0")

    report_message = None

    # =====================================================
    # USERS — NEW
    # =====================================================

    if report_type == "new_users":

        users = User.objects.filter(
            is_staff=False
        ).order_by(
            "-date_joined"
        )

        if start_date:

            users = users.filter(
                date_joined__date__gte=start_date,
                date_joined__date__lte=end_date,
            )

        report_columns = [
            "Username",
            "Name",
            "Email",
            "Joined",
        ]

        for user in users:

            report_rows.append({
                "username":
                    user.username,

                "name":
                    user.get_full_name()
                    or "-",

                "email":
                    user.email
                    or "-",

                "date":
                    user.date_joined,
            })

        report_count = len(report_rows)

    # =====================================================
    # USERS — ACTIVE
    # =====================================================

    elif report_type == "active_users":

        users = User.objects.filter(
            is_staff=False,
            last_login__isnull=False,
        ).order_by(
            "-last_login"
        )

        if start_date:

            users = users.filter(
                last_login__date__gte=start_date,
                last_login__date__lte=end_date,
            )

        report_columns = [
            "Username",
            "Name",
            "Email",
            "Last Active",
        ]

        for user in users:

            report_rows.append({
                "username":
                    user.username,

                "name":
                    user.get_full_name()
                    or "-",

                "email":
                    user.email
                    or "-",

                "date":
                    user.last_login,
            })

        report_count = len(report_rows)

    # =====================================================
    # USERS — INACTIVE
    # =====================================================

    elif report_type == "inactive_users":

        users = User.objects.filter(
            is_staff=False
        ).order_by(
            "last_login"
        )

        if start_date:

            users = users.filter(
                Q(last_login__isnull=True)
                |
                Q(
                    last_login__date__lt=
                    start_date
                )
            )

        report_columns = [
            "Username",
            "Name",
            "Email",
            "Last Active",
        ]

        for user in users:

            report_rows.append({
                "username":
                    user.username,

                "name":
                    user.get_full_name()
                    or "-",

                "email":
                    user.email
                    or "-",

                "date":
                    user.last_login,
            })

        report_count = len(report_rows)

    # =====================================================
    # BOOKINGS
    # =====================================================

    elif report_type in [
        "bookings",
        "accepted_bookings",
        "completed_bookings",
        "cancelled_bookings",
        "pending_bookings",
    ]:

        bookings = Booking.objects.select_related(
            "worker",
            "customer",
        ).order_by(
            "-created_at"
        )

        if start_date:

            bookings = bookings.filter(
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )

        status_map = {

            "accepted_bookings":
                "Accepted",

            "completed_bookings":
                "Completed",

            "cancelled_bookings":
                "Cancelled",

            "pending_bookings":
                "Pending",
        }

        if report_type in status_map:

            bookings = bookings.filter(
                status=status_map[report_type]
            )

        report_columns = [
            "Booking ID",
            "Date",
            "Customer",
            "Worker",
            "Work",
            "Status",
            "Amount",
        ]

        for booking in bookings:

            amount = (
                booking.final_amount
                if booking.final_amount is not None
                else Decimal("0")
            )

            report_rows.append({
                "id":
                    booking.id,

                "date":
                    booking.created_at,

                "customer":
                    booking.customer_name
                    or (
                        booking.customer.get_full_name()
                        if booking.customer
                        else "-"
                    ),

                "worker":
                    booking.worker.name
                    if booking.worker
                    else "-",

                "work":
                    booking.work_description
                    or "-",

                "status":
                    booking.status,

                "amount":
                    amount,
            })

            report_total += amount

        report_count = len(report_rows)

    # =====================================================
    # NEW WORKERS
    # =====================================================

    elif report_type == "new_workers":

        workers = Worker.objects.select_related(
            "user"
        ).order_by(
            "-user__date_joined"
        )

        if start_date:

            workers = workers.filter(
                user__date_joined__date__gte=start_date,
                user__date_joined__date__lte=end_date,
            )

        report_columns = [
            "Worker",
            "Profession",
            "Mobile",
            "City",
            "Experience",
            "Daily Wage",
            "Joined",
        ]

        for worker in workers:

            report_rows.append({
                "name":
                    worker.name,

                "profession":
                    worker.profession
                    or "-",

                "mobile":
                    worker.mobile
                    or "-",

                "city":
                    worker.city
                    or "-",

                "experience":
                    worker.experience
                    or "-",

                "daily_wage":
                    worker.daily_wage,

                "date":
                    worker.user.date_joined
                    if worker.user
                    else None,
            })

        report_count = len(report_rows)

    # =====================================================
    # REVIEWS
    # =====================================================

    elif report_type in [
        "reviews",
        "low_ratings",
    ]:

        reviews = Review.objects.select_related(
            "booking",
            "worker",
            "customer",
        ).order_by(
            "-created_at"
        )

        if start_date:

            reviews = reviews.filter(
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )

        if report_type == "low_ratings":

            reviews = reviews.filter(
                rating__lte=2
            )

        report_columns = [
            "Date",
            "Customer",
            "Worker",
            "Rating",
            "Comment",
            "Booking ID",
        ]

        for review in reviews:

            report_rows.append({
                "date":
                    review.created_at,

                "customer":
                    review.customer.get_full_name()
                    or review.customer.username,

                "worker":
                    review.worker.name
                    if review.worker
                    else "-",

                "rating":
                    review.rating,

                "comment":
                    review.comment
                    or "-",

                "booking":
                    review.booking.id
                    if review.booking
                    else "-",
            })

        report_count = len(report_rows)

    # =====================================================
    # FINANCIAL REPORTS
    # =====================================================

    elif report_type in [
        "platform_fees",
        "worker_payments",
        "outstanding",
        "ledger",
    ]:

        # -------------------------------------------------
        # SUPER ADMIN SECURITY
        # -------------------------------------------------

        if not request.user.is_superuser:

            raise PermissionDenied

        # -------------------------------------------------
        # LEDGER
        # -------------------------------------------------

        ledger = WorkerLedger.objects.select_related(
            "worker",
            "booking",
        ).order_by(
            "-created_at"
        )

        if start_date:

            ledger = ledger.filter(
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )

        if report_type == "platform_fees":

            ledger = ledger.filter(
                transaction_type="Booking Fee"
            )

        elif report_type == "worker_payments":

            ledger = ledger.filter(
                transaction_type="Payment"
            )

        elif report_type == "ledger":

            pass

        elif report_type == "outstanding":

            workers = Worker.objects.annotate(
                outstanding=Sum(
                    "ledger_entries__amount",
                    filter=Q(
                        ledger_entries__transaction_type=
                        "Booking Fee"
                    )
                    & Q(
                        ledger_entries__status="Pending"
                    )
                )
            ).order_by(
                "-outstanding"
            )

            report_columns = [
                "Worker",
                "Mobile",
                "Outstanding",
            ]

            for worker in workers:

                outstanding = (
                    worker.outstanding
                    or Decimal("0")
                )

                if outstanding > 0:

                    report_rows.append({
                        "worker":
                            worker.name,

                        "mobile":
                            worker.mobile
                            or "-",

                        "amount":
                            outstanding,
                    })

                    report_total += outstanding

            report_count = len(report_rows)

        if report_type != "outstanding":

            report_columns = [
                "Date",
                "Worker",
                "Transaction",
                "Amount",
                "Status",
                "Booking ID",
            ]

            for entry in ledger:

                amount = (
                    entry.amount
                    or Decimal("0")
                )

                report_rows.append({
                    "date":
                        entry.created_at,

                    "worker":
                        entry.worker.name
                        if entry.worker
                        else "-",

                    "transaction":
                        entry.transaction_type,

                    "amount":
                        amount,

                    "status":
                        entry.status
                        or "-",

                    "booking":
                        entry.booking.id
                        if entry.booking
                        else "-",
                })

                report_total += amount

            report_count = len(report_rows)

    # =====================================================
    # AUDIT LOGS
    # =====================================================

    elif report_type == "audit_logs":

        audit_logs = AuditLog.objects.select_related(
            "admin",
            "target_user",
        ).order_by(
            "-created_at"
        )

        if start_date:
            audit_logs = audit_logs.filter(
                created_at__date__gte=start_date,
                created_at__date__lte=end_date,
            )

        report_columns = [
            "Date & Time",
            "Admin",
            "Action",
            "Module",
            "Description",
            "Target User",
            "Target ID",
            "IP Address",
        ]

        for log in audit_logs:
            report_rows.append({
                "date":
                    log.created_at,

                "admin":
                    log.admin.username
                    if log.admin
                    else "System",

                "action":
                    log.get_action_display(),

                "module":
                    log.module,

                "description":
                    log.description
                    or "-",

                "target_user":
                    log.target_user.username
                    if log.target_user
                    else "-",

                "target_id":
                    log.target_id
                    or "-",

                "ip_address":
                    log.ip_address
                    or "-",
            })

        report_count = len(report_rows)
        
    # =====================================================
    # APP RETENTION
    # =====================================================

    elif report_type == "app_retention":

        report_message = (
            "Detailed app installation and uninstall "
            "tracking will become available after the "
            "Android app installation-tracking system "
            "is implemented."
        )

        report_columns = [
            "Metric",
            "Status",
        ]

        report_rows = [
            {
                "metric":
                    "App Installations",

                "status":
                    "Android tracking required",
            },
            {
                "metric":
                    "App Uninstalls",

                "status":
                    "Android tracking required",
            },
            {
                "metric":
                    "Reinstalls",

                "status":
                    "Android tracking required",
            },
            {
                "metric":
                    "Retention Rate",

                "status":
                    "Android tracking required",
            },
        ]

        report_count = 0

    # =====================================================
    # REPORT CONTEXT
    # =====================================================

    context = {

        "report_definitions":
            report_definitions,

        "selected_report":
            selected_report,

        "report_type":
            report_type,

        "period":
            period,

        "period_label":
            period_label,

        "start_date":
            start_date,

        "end_date":
            end_date,

        "custom_start":
            custom_start,

        "custom_end":
            custom_end,

        "report_columns":
            report_columns,

        "report_rows":
            report_rows,

        "report_count":
            report_count,

        "report_total":
            report_total,

        "report_message":
            report_message,

        "is_super_admin":
            request.user.is_superuser,
    }

    response = render(
        request,
        "admin_dashboard/reports.html",
        context,
    )

    # Keep the generated report data available
    # for Excel/PDF export views.
    response.report_context = context

    return response

# =========================================================
# REPORT EXPORT — EXCEL
# =========================================================

@admin_required
def export_report_excel(request):
    """
    Export the currently selected report to Excel.

    Uses the same report_type and date filters
    used by the Reports page.
    """

    report_response = admin_reports(request)

    context = getattr(
        report_response,
        "report_context",
        None
    )

    if not context:
        raise PermissionDenied(
            "Unable to generate report data."
        )

    report_type = context["report_type"]
    selected_report = context["selected_report"]

    report_columns = context["report_columns"]
    report_rows = context["report_rows"]

    period_label = context["period_label"]

    report_count = context["report_count"]
    report_total = context["report_total"]

    # -----------------------------------------------------
    # WORKBOOK
    # -----------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Report"

    # -----------------------------------------------------
    # TITLE
    # -----------------------------------------------------

    worksheet["A1"] = "KaamSetu"
    worksheet["A1"].font = Font(
        bold=True,
        size=18
    )

    worksheet["A2"] = selected_report["label"]
    worksheet["A2"].font = Font(
        bold=True,
        size=14
    )

    worksheet["A3"] = (
        f"Period: {period_label}"
    )

    worksheet["A4"] = (
        f"Total Records: {report_count}"
    )

    # -----------------------------------------------------
    # HEADERS
    # -----------------------------------------------------

    header_row = 6

    for column_number, column_name in enumerate(
        report_columns,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=column_name
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="17375E"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # -----------------------------------------------------
    # DATA
    # -----------------------------------------------------

    data_start_row = header_row + 1

    # Map report columns to actual row_data keys
    column_keys = {
        "bookings": [
            "id",
            "date",
            "customer",
            "worker",
            "work",
            "status",
            "amount",
        ],

        "accepted_bookings": [
            "id",
            "date",
            "customer",
            "worker",
            "work",
            "status",
            "amount",
        ],

        "completed_bookings": [
            "id",
            "date",
            "customer",
            "worker",
            "work",
            "status",
            "amount",
        ],

        "cancelled_bookings": [
            "id",
            "date",
            "customer",
            "worker",
            "work",
            "status",
            "amount",
        ],

        "pending_bookings": [
            "id",
            "date",
            "customer",
            "worker",
            "work",
            "status",
            "amount",
        ],
    }

    keys = column_keys.get(
        report_type,
        list(report_rows[0].keys())
        if report_rows
        else []
    )

    for row_number, row_data in enumerate(
        report_rows,
        start=data_start_row
    ):

        for column_number, key in enumerate(
            keys,
            start=1
        ):

            value = row_data.get(
                key,
                ""
            )

            # -------------------------------------------------
            # TIMEZONE-AWARE DATETIME
            # -------------------------------------------------

            if isinstance(value, datetime):

                if value.tzinfo is not None:

                    value = value.replace(
                        tzinfo=None
                    )

            # -------------------------------------------------
            # CREATE CELL
            # -------------------------------------------------

            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

            # -------------------------------------------------
            # DECIMAL → FLOAT
            # -------------------------------------------------

            if isinstance(value, Decimal):

                cell.value = float(value)

                cell.number_format = (
                    '₹#,##0.00'
                )

            # -------------------------------------------------
            # DATETIME FORMATTING
            # -------------------------------------------------

            elif isinstance(value, datetime):

                cell.number_format = (
                    "dd-mmm-yyyy hh:mm AM/PM"
                )

            cell.alignment = Alignment(
                vertical="top"
            )

    # -----------------------------------------------------
    # TOTAL
    # -----------------------------------------------------

    if report_total:

        total_row = (
            data_start_row
            + len(report_rows)
            + 1
        )

        worksheet.cell(
            row=total_row,
            column=max(
                1,
                len(report_columns) - 1
            ),
            value="Total"
        ).font = Font(
            bold=True
        )

        total_cell = worksheet.cell(
            row=total_row,
            column=len(report_columns),
            value=float(report_total)
        )

        total_cell.font = Font(
            bold=True
        )

        total_cell.number_format = (
            '₹#,##0.00'
        )

    # -----------------------------------------------------
    # COLUMN WIDTH
    # -----------------------------------------------------

    for column_cells in worksheet.columns:

        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = 0

        for cell in column_cells:

            if cell.value is not None:

                value_length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    value_length
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 3, 12),
            45
        )

    # -----------------------------------------------------
    # FREEZE HEADER
    # -----------------------------------------------------

    worksheet.freeze_panes = "A7"

    # -----------------------------------------------------
    # AUTO FILTER
    # -----------------------------------------------------

    if report_columns:

        last_column = get_column_letter(
            len(report_columns)
        )

        last_row = (
            data_start_row
            + len(report_rows)
            - 1
        )

        if last_row >= header_row:

            worksheet.auto_filter.ref = (
                f"A{header_row}:"
                f"{last_column}{last_row}"
            )

    # -----------------------------------------------------
    # RESPONSE
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"KaamSetu_"
        f"{report_type}_"
        f"report.xlsx"
    )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response



# =========================================================
# REPORT EXPORT — PDF
# =========================================================

@admin_required
def export_report_pdf(request):

    """
    Export the currently selected report to PDF.
    """

    report_response = admin_reports(request)

    context = getattr(
        report_response,
        "report_context",
        None
    )

    if not context:
        raise PermissionDenied(
            "Unable to generate report data."
        )

    report_type = context["report_type"]

    selected_report = context["selected_report"]

    report_columns = context["report_columns"]

    report_rows = context["report_rows"]

    period_label = context["period_label"]

    report_count = context["report_count"]

    report_total = context["report_total"]

    # -----------------------------------------------------
    # PDF BUFFER
    # -----------------------------------------------------

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "KaamSetuTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#17375E"),
        spaceAfter=6,
    )

    report_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#17375E"),
        spaceAfter=5,
    )

    info_style = ParagraphStyle(
        "ReportInfo",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#5F6B7A"),
        spaceAfter=3,
    )

    cell_style = ParagraphStyle(
        "Cell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
    )

    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    story = []

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    story.append(
        Paragraph(
            "KaamSetu",
            title_style
        )
    )

    story.append(
        Paragraph(
            "Admin Report",
            report_style
        )
    )

    story.append(
        Paragraph(
            f"<b>Report:</b> "
            f"{selected_report['label']}",
            info_style
        )
    )

    story.append(
        Paragraph(
            f"<b>Period:</b> "
            f"{period_label}",
            info_style
        )
    )

    story.append(
        Paragraph(
            f"<b>Total Records:</b> "
            f"{report_count}",
            info_style
        )
    )

    if report_total:

        story.append(
            Paragraph(
                f"<b>Total Amount:</b> "
                f"₹{report_total:,.2f}",
                info_style
            )
        )

    story.append(
        Spacer(1, 8)
    )

    # -----------------------------------------------------
    # TABLE
    # -----------------------------------------------------

    table_data = []

    table_data.append([
        Paragraph(
            str(column),
            header_cell_style
        )
        for column in report_columns
    ])

    for row_data in report_rows:

        values = list(
            row_data.values()
        )

        formatted_values = []

        for value in values:

            if isinstance(value, Decimal):

                display_value = (
                    f"₹{value:,.2f}"
                )

            elif hasattr(value, "strftime"):

                display_value = value.strftime(
                    "%d %b %Y, %I:%M %p"
                )

            else:

                display_value = (
                    "-"
                    if value is None
                    else str(value)
                )

            formatted_values.append(
                Paragraph(
                    display_value,
                    cell_style
                )
            )

        table_data.append(
            formatted_values
        )

    # -----------------------------------------------------
    # EMPTY REPORT
    # -----------------------------------------------------

    if len(table_data) == 1:

        table_data.append([
            Paragraph(
                "No records found for the selected report and period.",
                cell_style
            )
        ])

        table = Table(
            table_data,
            colWidths=[260 * mm]
        )

    else:

        available_width = 273 * mm

        column_count = len(
            report_columns
        )

        column_width = (
            available_width /
            max(column_count, 1)
        )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[
                column_width
            ] * column_count
        )

    # -----------------------------------------------------
    # TABLE STYLE
    # -----------------------------------------------------

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#17375E")
            ),

            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),

            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),

            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "TOP"
            ),

            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.35,
                colors.HexColor("#D9DEE5")
            ),

            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#F7F9FC"),
                ]
            ),

            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                5
            ),

            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                5
            ),

            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5
            ),

            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5
            ),
        ])
    )

    story.append(table)

    # -----------------------------------------------------
    # BUILD
    # -----------------------------------------------------

    document.build(story)

    buffer.seek(0)

    filename = (
        f"KaamSetu_"
        f"{report_type}_"
        f"report.pdf"
    )

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    return response

# =========================================================
# SUPER ADMIN ONLY
# ADMIN MANAGEMENT SECURITY
# =========================================================

def super_admin_required(view_func):
    """
    Allow ONLY Django Super Admin users.

    Normal Admin users are explicitly denied.
    """

    @wraps(view_func)
    def wrapper(request, *args, **kwargs):

        if not request.user.is_authenticated:
            return redirect("login")

        if not request.user.is_superuser:
            raise PermissionDenied(
                "Only Super Admin can access Admin Management."
            )

        return view_func(
            request,
            *args,
            **kwargs
        )

    return wrapper


# =========================================================
# ADMIN MANAGEMENT
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_management(request):

    # -----------------------------------------------------
    # ONLY NORMAL ADMINS
    #
    # Super Admin itself is not included in this list.
    # -----------------------------------------------------

    admins = User.objects.filter(
        is_staff=True,
        is_superuser=False
    ).prefetch_related(
        "groups",
        "user_permissions"
    ).order_by(
        "username"
    )

    # -----------------------------------------------------
    # ALL AVAILABLE GROUPS
    # -----------------------------------------------------

    groups = Group.objects.all().order_by(
        "name"
    )

    return render(
        request,
        "admin_dashboard/admin_management.html",
        {
            "admins": admins,
            "groups": groups,
        }
    )
    
# =========================================================
# AUDIT LOGS
# ADMIN ACCESS
# =========================================================

@admin_required
def admin_audit_logs(request):

    audit_logs = AuditLog.objects.select_related(
        "admin",
        "target_user",
    ).all()

    return render(
        request,
        "admin_dashboard/audit_logs.html",
        {
            "audit_logs": audit_logs,
        }
    )

# =========================================================
# SYSTEM SETTINGS
# SUPER ADMIN ONLY
# =========================================================

@super_admin_required
def admin_system_settings(request):

    system_settings = SystemSetting.objects.first()

    # Create default settings if none exist
    if system_settings is None:
        system_settings = SystemSetting.objects.create(
            maintenance_mode=False,
            maintenance_message=(
                "KaamSetu is currently under maintenance. "
                "Please check back soon."
            ),
            support_email="",
            support_phone="",
            app_version="1.0.0",
        )

    if request.method == "POST":

        maintenance_mode = (
            request.POST.get("maintenance_mode") == "on"
        )

        maintenance_message = (
            request.POST.get(
                "maintenance_message",
                ""
            ).strip()
        )

        support_email = (
            request.POST.get(
                "support_email",
                ""
            ).strip()
        )

        support_phone = (
            request.POST.get(
                "support_phone",
                ""
            ).strip()
        )

        app_version = (
            request.POST.get(
                "app_version",
                ""
            ).strip()
        )

        # -------------------------------------------------
        # VALIDATION
        # -------------------------------------------------

        if not app_version:
            messages.error(
                request,
                "Application version is required."
            )
            return redirect("admin_system_settings")

        if maintenance_mode and not maintenance_message:
            messages.error(
                request,
                "Please enter a maintenance message."
            )
            return redirect("admin_system_settings")

        # -------------------------------------------------
        # SAVE
        # -------------------------------------------------

        system_settings.maintenance_mode = maintenance_mode
        system_settings.maintenance_message = maintenance_message
        system_settings.support_email = support_email
        system_settings.support_phone = support_phone
        system_settings.app_version = app_version
        system_settings.updated_by = request.user

        system_settings.save()

        # -------------------------------------------------
        # AUDIT LOG
        # -------------------------------------------------

        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="System Settings",
            description=(
                "System Settings updated. "
                f"Maintenance Mode: "
                f"{'ON' if maintenance_mode else 'OFF'}, "
                f"App Version: {app_version}."
            ),
            target_id=system_settings.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "System Settings updated successfully."
        )

        return redirect("admin_system_settings")

    return render(
        request,
        "admin_dashboard/system_settings.html",
        {
            "system_settings": system_settings,
        }
    )

@super_admin_required
def change_super_admin_password(request):
    form = PasswordChangeForm(
        user=request.user,
        data=request.POST or None
    )

    if request.method == "POST" and form.is_valid():
        user = form.save()

        # Keep Super Admin logged in after password change
        update_session_auth_hash(request, user)

        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="Password Management",
            description="Super Admin password changed successfully.",
            target_user=request.user,
            target_id=request.user.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Super Admin password changed successfully."
        )

        return redirect("change_super_admin_password")

    return render(
        request,
        "admin_dashboard/change_super_admin_password.html",
        {
            "form": form,
        }
    )

# =========================================================
# BOOKING EXPIRY
# SUPER ADMIN ONLY - MANUAL TRIGGER
# =========================================================

@super_admin_required
def expire_pending_bookings_manual(request):

    if request.method != "POST":
        raise PermissionDenied(
            "Invalid request method."
        )

    expired_count = expire_pending_bookings()

    # -----------------------------------------------------
    # AUDIT LOG
    # -----------------------------------------------------

    create_audit_log(
        admin=request.user,
        action="UPDATE",
        module="Booking Expiry",
        description=(
            f"Manual pending booking expiry executed. "
            f"{expired_count} booking(s) expired."
        ),
        ip_address=get_client_ip(request),
    )

    # -----------------------------------------------------
    # SUCCESS MESSAGE
    # -----------------------------------------------------

    if expired_count:
        messages.success(
            request,
            f"{expired_count} pending booking(s) "
            "expired successfully."
        )
    else:
        messages.info(
            request,
            "No pending booking(s) older than 1 hour found."
        )

    return redirect("admin_dashboard")

# =========================================================
# CAREERS / VOLUNTEER INTERNSHIP
# ADMIN LEVEL MANAGEMENT
# =========================================================

from careers.models import Vacancy, Application


@permission_required("careers.view_vacancy")
def admin_careers(request):

    vacancies = Vacancy.objects.order_by(
        "-created_at"
    )

    return render(
        request,
        "admin_dashboard/careers.html",
        {
            "vacancies": vacancies,
        }
    )


@permission_required("careers.add_vacancy")
def admin_career_create(request):

    if request.method == "POST":

        title = request.POST.get("title", "").strip()
        short_description = request.POST.get(
            "short_description", ""
        ).strip()
        description = request.POST.get(
            "description", ""
        ).strip()
        eligibility = request.POST.get(
            "eligibility", ""
        ).strip()
        location = request.POST.get(
            "location", ""
        ).strip()
        internship_type = request.POST.get(
            "internship_type",
            "Volunteer Internship"
        ).strip()
        stipend = request.POST.get(
            "stipend", ""
        ).strip()
        openings = request.POST.get(
            "openings", "1"
        ).strip()
        application_deadline = request.POST.get(
            "application_deadline"
        ) or None
        is_active = request.POST.get(
            "is_active"
        ) == "on"

        if not title or not short_description or not description:
            messages.error(
                request,
                "Title, short description and description are required."
            )
            return redirect("admin_career_create")

        try:
            openings = int(openings)

            if openings < 1:
                raise ValueError

        except (TypeError, ValueError):
            messages.error(
                request,
                "Openings must be a valid number greater than 0."
            )
            return redirect("admin_career_create")

        vacancy = Vacancy.objects.create(
            title=title,
            short_description=short_description,
            description=description,
            eligibility=eligibility,
            location=location,
            internship_type=internship_type or "Volunteer Internship",
            stipend=stipend,
            openings=openings,
            application_deadline=application_deadline,
            is_active=is_active,
        )

        create_audit_log(
            admin=request.user,
            action="CREATE",
            module="Careers",
            description=(
                f'Career opportunity "{vacancy.title}" created.'
            ),
            target_id=vacancy.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Career opportunity created successfully."
        )

        return redirect("admin_careers")

    return render(
        request,
        "admin_dashboard/career_form.html",
        {
            "page_title": "Create Career Opportunity",
            "form_action": "admin_career_create",
            "vacancy": None,
        }
    )


@permission_required("careers.change_vacancy")
def admin_career_edit(request, vacancy_id):

    vacancy = get_object_or_404(
        Vacancy,
        id=vacancy_id
    )

    if request.method == "POST":

        vacancy.title = request.POST.get(
            "title", ""
        ).strip()

        vacancy.short_description = request.POST.get(
            "short_description", ""
        ).strip()

        vacancy.description = request.POST.get(
            "description", ""
        ).strip()

        vacancy.eligibility = request.POST.get(
            "eligibility", ""
        ).strip()

        vacancy.location = request.POST.get(
            "location", ""
        ).strip()

        vacancy.internship_type = request.POST.get(
            "internship_type",
            "Volunteer Internship"
        ).strip() or "Volunteer Internship"

        vacancy.stipend = request.POST.get(
            "stipend", ""
        ).strip()

        openings_raw = request.POST.get(
            "openings", "1"
        ).strip()

        vacancy.application_deadline = (
            request.POST.get("application_deadline")
            or None
        )

        vacancy.is_active = (
            request.POST.get("is_active") == "on"
        )

        try:
            vacancy.openings = int(openings_raw)

            if vacancy.openings < 1:
                raise ValueError

        except (TypeError, ValueError):
            messages.error(
                request,
                "Openings must be a valid number greater than 0."
            )
            return redirect(
                "admin_career_edit",
                vacancy_id=vacancy.id
            )

        vacancy.save()

        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="Careers",
            description=(
                f'Career opportunity "{vacancy.title}" updated.'
            ),
            target_id=vacancy.id,
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Career opportunity updated successfully."
        )

        return redirect("admin_careers")

    return render(
        request,
        "admin_dashboard/career_form.html",
        {
            "page_title": "Edit Career Opportunity",
            "form_action": "admin_career_edit",
            "vacancy": vacancy,
        }
    )


@permission_required("careers.change_vacancy")
@require_POST
def admin_career_toggle(request, vacancy_id):

    vacancy = get_object_or_404(
        Vacancy,
        id=vacancy_id
    )

    vacancy.is_active = not vacancy.is_active

    vacancy.save(
        update_fields=[
            "is_active",
            "updated_at",
        ]
    )

    state = (
        "activated"
        if vacancy.is_active
        else "deactivated"
    )

    create_audit_log(
        admin=request.user,
        action="UPDATE",
        module="Careers",
        description=(
            f'Career opportunity "{vacancy.title}" '
            f'was {state}.'
        ),
        target_id=vacancy.id,
        ip_address=get_client_ip(request),
    )

    messages.success(
        request,
        f'Career opportunity "{vacancy.title}" '
        f'{state} successfully.'
    )

    return redirect("admin_careers")

# =========================================================
# CAREERS / VOLUNTEER INTERNSHIP
# APPLICATION MANAGEMENT
# =========================================================

from careers.models import Application


@permission_required("careers.view_application")
def admin_career_applications(request):
    applications = Application.objects.select_related(
        "vacancy"
    ).order_by("-applied_at")

    return render(
        request,
        "admin_dashboard/career_applications.html",
        {
            "applications": applications,
        },
    )


@permission_required("careers.change_application")
def admin_career_application_edit(request, application_id):
    application = get_object_or_404(
        Application,
        id=application_id,
    )

    if request.method == "POST":
        status = request.POST.get("status", "").strip()
        admin_notes = request.POST.get("admin_notes", "").strip()
        interview_date = request.POST.get("interview_date", "").strip()
        interview_notes = request.POST.get("interview_notes", "").strip()

        certificate_eligible = (
            request.POST.get("certificate_eligible") == "on"
        )

        certificate_issued = (
            request.POST.get("certificate_issued") == "on"
        )

        valid_statuses = {
            choice[0]
            for choice in Application.STATUS_CHOICES
        }

        if status not in valid_statuses:
            messages.error(
                request,
                "Invalid application status."
            )
            return redirect(
                "admin_career_application_edit",
                application_id=application.id,
            )

        application.status = status
        application.admin_notes = admin_notes
        application.interview_notes = interview_notes
        application.certificate_eligible = certificate_eligible
        application.certificate_issued = certificate_issued

        if interview_date:
            try:
                application.interview_date = datetime.fromisoformat(
                    interview_date
                )
            except ValueError:
                messages.error(
                    request,
                    "Invalid interview date/time."
                )
                return redirect(
                    "admin_career_application_edit",
                    application_id=application.id,
                )
        else:
            application.interview_date = None

        application.save()

        create_audit_log(
            admin=request.user,
            action="UPDATE",
            module="Careers",
            description=(
                f"Career application updated: "
                f"{application.full_name} - "
                f"{application.vacancy.title}. "
                f"Status: {application.status}."
            ),
            ip_address=get_client_ip(request),
        )

        messages.success(
            request,
            "Career application updated successfully."
        )

        return redirect(
            "admin_career_applications"
        )

    return render(
        request,
        "admin_dashboard/career_application_form.html",
        {
            "application": application,
        },
    )